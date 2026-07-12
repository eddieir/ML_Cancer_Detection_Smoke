"""
Builds a full demo report (DOCX + PDF) from the logs produced by
`scripts/run_demo.sh`: pytest results, each module's smoke-test output, and
the evaluation metrics JSON produced by `src/evaluate.py`.

Usage: python3 scripts/generate_report.py --run-dir demo_run_<timestamp>
"""
import argparse
import json
import re
import subprocess
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).parents[1]

STEPS = [
    ("00_pip_install.log",        "Install project dependencies"),
    ("02_pytest.log",             "Automated test suite (pytest)"),
    ("03_preprocess.log",         "preprocess.py — data pipeline smoke test"),
    ("04_model.log",              "model.py — MultiSmokeCancerNet smoke test"),
    ("05_train.log",              "train.py — 3-phase training smoke test"),
    ("06_evaluate.log",           "evaluate.py — evaluation + interpretability smoke test"),
    ("07_inference.log",          "inference.py — prediction smoke test"),
    ("08_generate_plots.log",     "generate_plots.py — render all figures"),
]

PYTEST_SUMMARY_RE = re.compile(r"^={2,}.*\d+ (passed|failed).*={2,}$", re.MULTILINE)

PLOT_CAPTIONS = {
    "training_history":          "Training curves across all three curriculum phases",
    "roc_curve":                 "Subject-level cancer risk — ROC curve",
    "pr_curve":                  "Subject-level cancer risk — precision-recall curve",
    "calibration_curve":         "Subject-level cancer risk — calibration curve",
    "confusion_matrix_smoke":    "Cell-level smoke-type classification — confusion matrix",
    "attention_by_cell_type":    "MIL attention weight, averaged by cell type",
    "attention_by_smoke_type":   "MIL attention weight, averaged by smoke type",
    "cell_umap":                 "UMAP of cell embeddings (smoke type and malignancy score)",
}


def read_log(run_dir: Path, filename: str) -> str:
    path = run_dir / "logs" / filename
    if not path.exists():
        return "(log not found — step did not run)"
    return path.read_text()


def step_passed(run_dir: Path, filename: str) -> bool:
    exit_path = run_dir / "logs" / f"{Path(filename).stem}.exit"
    if exit_path.exists():
        return exit_path.read_text().strip() == "0"
    return False


def pytest_summary_line(log_text: str) -> str:
    m = PYTEST_SUMMARY_RE.search(log_text)
    return m.group(0).strip("= ").strip() if m else "no summary line found"


def load_novel_contributions():
    """
    Parses the "Novel Contributions vs. Literature" markdown table straight
    out of ARCHITECTURE.md, so the report and the architecture doc never
    drift out of sync.
    """
    path = ROOT / "ARCHITECTURE.md"
    if not path.exists():
        return None
    text = path.read_text()
    m = re.search(
        r"^## \d+\. Novel Contributions vs\. Literature\s*\n(.*?)(?=\n## |\Z)",
        text, re.DOTALL | re.MULTILINE,
    )
    if not m:
        return None
    lines = [ln.strip() for ln in m.group(1).strip().splitlines() if ln.strip().startswith("|")]
    if len(lines) < 2:
        return None
    rows = []
    for ln in lines:
        if set(ln.replace("|", "").strip()) <= set("-: "):
            continue  # markdown header separator row
        cells = [c.strip() for c in ln.strip("|").split("|")]
        rows.append(cells)
    return rows  # rows[0] is the header


def collect_plots(plots_dir: Path):
    if not plots_dir or not plots_dir.exists():
        return []
    plots = []
    for path in sorted(plots_dir.glob("*.png")):
        caption = PLOT_CAPTIONS.get(path.stem)
        if caption is None:
            caption = path.stem.replace("_", " ").capitalize()
        plots.append({"path": path, "caption": caption})
    return plots


def load_eval_report() -> dict:
    path = ROOT / "checkpoints" / "evaluation_report.json"
    if not path.exists():
        return {}
    return json.loads(path.read_text())


def git_info() -> dict:
    def run(*args):
        try:
            return subprocess.check_output(["git", *args], cwd=ROOT, text=True).strip()
        except Exception:
            return "unknown"
    return {
        "branch": run("rev-parse", "--abbrev-ref", "HEAD"),
        "commit": run("rev-parse", "--short", "HEAD"),
    }


def truncate(text: str, max_lines: int = 60) -> str:
    lines = text.splitlines()
    if len(lines) <= max_lines:
        return text
    head = lines[: max_lines // 2]
    tail = lines[-max_lines // 2:]
    return "\n".join(head + [f"... ({len(lines) - max_lines} lines omitted) ..."] + tail)


def build_sections(run_dir: Path):
    sections = []
    for filename, title in STEPS:
        text = read_log(run_dir, filename)
        passed = step_passed(run_dir, filename)
        sections.append({
            "title": title,
            "passed": passed,
            "summary": pytest_summary_line(text) if "pytest" in filename else None,
            "log": truncate(text),
        })
    return sections


def build_docx(run_dir: Path, sections, eval_report, plots, novel_contributions, info, out_path: Path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    title = doc.add_heading("MultiSmokeCancerNet — Demo Run Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Generated {info['generated_at']}  |  branch {info['branch']}  |  commit {info['commit']}"
    ).italic = True

    doc.add_heading("Overview", level=1)
    doc.add_paragraph(
        "This report documents an automated run of the MultiSmokeCancerNet "
        "pipeline: the pytest suite, and every module's built-in synthetic-"
        "data smoke test (preprocessing, model, training, evaluation, "
        "inference). All steps run on generated synthetic data with no "
        "external network downloads, so this verifies the pipeline is wired "
        "together correctly end to end."
    )

    doc.add_heading("Results at a glance", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Step", "Result"
    for sec in sections:
        row = table.add_row().cells
        row[0].text = sec["title"]
        row[1].text = "PASSED" if sec["passed"] else "FAILED"
        color = RGBColor(0x1a, 0x7f, 0x37) if sec["passed"] else RGBColor(0xcf, 0x22, 0x2e)
        for run in row[1].paragraphs[0].runs:
            run.font.color.rgb = color
            run.font.bold = True

    if eval_report:
        doc.add_heading("Evaluation metrics (smoke test, untrained model on random data)", level=1)
        doc.add_paragraph(
            "These numbers come from evaluate.py's smoke test, which runs an "
            "untrained model on random synthetic data purely to check the "
            "evaluation code path. Near-random values (macro-F1 ~ 0.04, "
            "AUC ~ 0.5) are expected and correct here — they are not a "
            "measure of model quality."
        )
        cl = eval_report.get("cell_level", {})
        sl = eval_report.get("subject_level", {})
        ip = eval_report.get("interpretability", {})
        metrics_table = doc.add_table(rows=1, cols=2)
        metrics_table.style = "Light Grid Accent 1"
        h = metrics_table.rows[0].cells
        h[0].text, h[1].text = "Metric", "Value"
        rows = [
            ("Cells evaluated", cl.get("n_cells")),
            ("Smoke-type macro F1", cl.get("smoke_type", {}).get("macro_f1")),
            ("Malignancy ROC-AUC (cell level)", cl.get("malignancy", {}).get("roc_auc")),
            ("Subjects evaluated", sl.get("n_subjects")),
            ("Cancer ROC-AUC (subject level)", sl.get("cancer", {}).get("roc_auc")),
            ("Sensitivity", sl.get("cancer", {}).get("sensitivity")),
            ("Specificity", sl.get("cancer", {}).get("specificity")),
            ("Malignancy-attention correlation", ip.get("malignancy_attention_correlation")),
        ]
        for name, value in rows:
            r = metrics_table.add_row().cells
            r[0].text = str(name)
            r[1].text = str(value)

    if plots:
        doc.add_heading("Figures", level=1)
        for plot in plots:
            doc.add_heading(plot["caption"], level=2)
            doc.add_picture(str(plot["path"]), width=Inches(6))

    if novel_contributions:
        doc.add_heading("Novel Contributions vs. Literature", level=1)
        doc.add_paragraph(
            "Reproduced from ARCHITECTURE.md — how this project's claims compare "
            "to the closest existing published work."
        )
        header, *rows = novel_contributions
        nc_table = doc.add_table(rows=1, cols=len(header))
        nc_table.style = "Light Grid Accent 1"
        for cell, text in zip(nc_table.rows[0].cells, header):
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in rows:
            cells = nc_table.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = text

    doc.add_heading("Step-by-step logs", level=1)
    for sec in sections:
        doc.add_heading(sec["title"], level=2)
        status_p = doc.add_paragraph()
        run = status_p.add_run("PASSED" if sec["passed"] else "FAILED")
        run.bold = True
        run.font.color.rgb = RGBColor(0x1a, 0x7f, 0x37) if sec["passed"] else RGBColor(0xcf, 0x22, 0x2e)
        if sec["summary"]:
            doc.add_paragraph(f"pytest summary: {sec['summary']}")
        log_p = doc.add_paragraph()
        log_run = log_p.add_run(sec["log"])
        log_run.font.name = "Courier New"
        log_run.font.size = Pt(8)

    doc.save(str(out_path))


def build_pdf(run_dir: Path, sections, eval_report, plots, novel_contributions, info, out_path: Path):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image,
    )

    styles = getSampleStyleSheet()
    mono = ParagraphStyle("mono", parent=styles["Code"], fontSize=6.5, leading=8)
    body = styles["BodyText"]

    doc = SimpleDocTemplate(str(out_path), pagesize=LETTER,
                             leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                             topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    story = []

    story.append(Paragraph("MultiSmokeCancerNet — Demo Run Report", styles["Title"]))
    story.append(Paragraph(
        f"Generated {info['generated_at']} | branch {info['branch']} | commit {info['commit']}",
        styles["Italic"]))
    story.append(Spacer(1, 0.2 * inch))

    story.append(Paragraph("Overview", styles["Heading1"]))
    story.append(Paragraph(
        "This report documents an automated run of the MultiSmokeCancerNet "
        "pipeline: the pytest suite, and every module's built-in synthetic-"
        "data smoke test (preprocessing, model, training, evaluation, "
        "inference). All steps run on generated synthetic data with no "
        "external network downloads, so this verifies the pipeline is wired "
        "together correctly end to end.", body))
    story.append(Spacer(1, 0.15 * inch))

    story.append(Paragraph("Results at a glance", styles["Heading1"]))
    data = [["Step", "Result"]] + [
        [sec["title"], "PASSED" if sec["passed"] else "FAILED"] for sec in sections
    ]
    t = Table(data, colWidths=[4.5 * inch, 1.2 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f4f6f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    for i, sec in enumerate(sections, start=1):
        color = colors.HexColor("#1a7f37") if sec["passed"] else colors.HexColor("#cf222e")
        t.setStyle(TableStyle([("TEXTCOLOR", (1, i), (1, i), color)]))
    story.append(t)
    story.append(Spacer(1, 0.2 * inch))

    if eval_report:
        story.append(Paragraph(
            "Evaluation metrics (smoke test, untrained model on random data)",
            styles["Heading1"]))
        story.append(Paragraph(
            "These numbers come from evaluate.py's smoke test, which runs an "
            "untrained model on random synthetic data purely to check the "
            "evaluation code path. Near-random values (macro-F1 ~ 0.04, "
            "AUC ~ 0.5) are expected and correct here.", body))
        cl = eval_report.get("cell_level", {})
        sl = eval_report.get("subject_level", {})
        ip = eval_report.get("interpretability", {})
        rows = [
            ("Cells evaluated", cl.get("n_cells")),
            ("Smoke-type macro F1", cl.get("smoke_type", {}).get("macro_f1")),
            ("Malignancy ROC-AUC (cell level)", cl.get("malignancy", {}).get("roc_auc")),
            ("Subjects evaluated", sl.get("n_subjects")),
            ("Cancer ROC-AUC (subject level)", sl.get("cancer", {}).get("roc_auc")),
            ("Sensitivity", sl.get("cancer", {}).get("sensitivity")),
            ("Specificity", sl.get("cancer", {}).get("specificity")),
            ("Malignancy-attention correlation", ip.get("malignancy_attention_correlation")),
        ]
        mdata = [["Metric", "Value"]] + [[str(n), str(v)] for n, v in rows]
        mt = Table(mdata, colWidths=[3.8 * inch, 1.9 * inch])
        mt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f4f6f")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(mt)

    if plots:
        story.append(PageBreak())
        story.append(Paragraph("Figures", styles["Heading1"]))
        max_width = 6.1 * inch
        for plot in plots:
            story.append(Paragraph(plot["caption"], styles["Heading2"]))
            reader = ImageReader(str(plot["path"]))
            img_w, img_h = reader.getSize()
            scale = min(max_width / img_w, 1.0)
            story.append(Image(str(plot["path"]), width=img_w * scale, height=img_h * scale))
            story.append(Spacer(1, 0.15 * inch))

    if novel_contributions:
        story.append(PageBreak())
        story.append(Paragraph("Novel Contributions vs. Literature", styles["Heading1"]))
        story.append(Paragraph(
            "Reproduced from ARCHITECTURE.md — how this project's claims compare "
            "to the closest existing published work.", body))
        story.append(Spacer(1, 0.1 * inch))
        cell_style = ParagraphStyle("nc_cell", parent=body, fontSize=7.5, leading=9)
        header, *rows = novel_contributions
        nc_data = [[Paragraph(f"<b>{c}</b>", cell_style) for c in header]]
        for row in rows:
            nc_data.append([Paragraph(c, cell_style) for c in row])
        nc_table = Table(nc_data, colWidths=[1.7 * inch, 1.9 * inch, 2.1 * inch])
        nc_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f4f6f")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(nc_table)

    story.append(PageBreak())
    story.append(Paragraph("Step-by-step logs", styles["Heading1"]))
    for sec in sections:
        story.append(Paragraph(sec["title"], styles["Heading2"]))
        status_color = "green" if sec["passed"] else "red"
        status_text = "PASSED" if sec["passed"] else "FAILED"
        story.append(Paragraph(f'<font color="{status_color}"><b>{status_text}</b></font>', body))
        if sec["summary"]:
            story.append(Paragraph(f"pytest summary: {sec['summary']}", body))
        escaped = (
            sec["log"]
            .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        )
        story.append(Paragraph(escaped.replace("\n", "<br/>"), mono))
        story.append(Spacer(1, 0.15 * inch))

    doc.build(story)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", required=True, type=Path)
    parser.add_argument("--plots-dir", type=Path, default=None)
    args = parser.parse_args()

    run_dir = args.run_dir
    sections = build_sections(run_dir)
    eval_report = load_eval_report()
    plots = collect_plots(args.plots_dir)
    novel_contributions = load_novel_contributions()
    info = {**git_info(), "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}

    docx_path = run_dir / "report.docx"
    pdf_path = run_dir / "report.pdf"

    build_docx(run_dir, sections, eval_report, plots, novel_contributions, info, docx_path)
    print(f"Wrote {docx_path}")

    build_pdf(run_dir, sections, eval_report, plots, novel_contributions, info, pdf_path)
    print(f"Wrote {pdf_path}")

    any_failed = any(not sec["passed"] for sec in sections)
    return 1 if any_failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
